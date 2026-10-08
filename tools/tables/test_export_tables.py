"""Run with: python -I tools/tables/test_export_tables.py"""
import json
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import export_tables  # noqa: E402


class ExportTablesTests(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="bk_tables_")
        self.sample = os.path.join(self.dir, "Sample.xlsx")
        export_tables.write_sample(self.sample)

    def test_sample_exports_json_and_csharp(self):
        json_out = os.path.join(self.dir, "json")
        cs_out = os.path.join(self.dir, "cs")
        written = export_tables.export([self.sample], json_out, cs_out, "Project.Tables")
        self.assertEqual(sorted(os.path.basename(w) for w in written), ["Stage.json", "StageTable.cs", "TableEnums.cs"])

        with open(os.path.join(json_out, "Stage.json"), encoding="utf-8") as f:
            data = json.load(f)
        rows = data["_rows"]
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0], {"id": 1, "reward": 100, "tags": ["a", "b"], "difficulty": 2})
        self.assertEqual(rows[1]["tags"], [])
        self.assertEqual(rows[1]["difficulty"], 0)
        self.assertNotIn("note", rows[0], "columns not marked 'data' are skipped")

        with open(os.path.join(cs_out, "StageTable.cs"), encoding="utf-8") as f:
            cs = f.read()
        self.assertIn("public sealed class StageRow : ITableRow<int>", cs)
        self.assertIn("public string[] tags;", cs)
        self.assertIn("public DIFFICULTY difficulty;", cs)
        self.assertIn("public int Id => id;", cs)
        self.assertIn("class StageTable : TableAsset<int, StageRow>", cs)
        with open(os.path.join(cs_out, "TableEnums.cs"), encoding="utf-8") as f:
            enums = f.read()
        self.assertIn("public enum DIFFICULTY", enums)
        self.assertIn("HARD = 2,", enums)

    def test_duplicate_key_fails(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.sample)
        wb["Stage"].cell(row=6, column=2, value=1)
        bad = os.path.join(self.dir, "Bad.xlsx")
        wb.save(bad)
        with self.assertRaises(export_tables.ExportError):
            export_tables.export([bad], os.path.join(self.dir, "j"), os.path.join(self.dir, "c"), "X")
        self.assertEqual(export_tables.main(["--xlsx", bad, "--json-out", os.path.join(self.dir, "j2"),
                                             "--cs-out", os.path.join(self.dir, "c2")]), 1)

    def test_sheet_without_data_marker_is_skipped(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.sample)
        wb["Stage"].cell(row=2, column=1, value="#note")
        other = os.path.join(self.dir, "Other.xlsx")
        wb.save(other)
        with self.assertRaises(export_tables.ExportError):
            export_tables.export([other], os.path.join(self.dir, "j"), os.path.join(self.dir, "c"), "X")


if __name__ == "__main__":
    unittest.main()
