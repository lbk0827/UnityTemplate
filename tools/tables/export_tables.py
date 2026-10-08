#!/usr/bin/env python3
"""Export xlsx tables to BK.Data JSON + C# (TableAsset<TKey, TRow>).

Header convention (compatible with the sf/RootBox exporter, 1-based sheet rows):
  row 1, col A : optional sub-folder (ignored here)
  row 2, col A : "#data"  (sheets without it are skipped)
  row 2        : column names; a trailing "[]" marks an array column
  row 3        : comma list of type/options per column:
                 key | int | long | float | double | bool | string | type<ENUM> | unique
                 (check<...>, localtext, stream, origin, null are accepted and ignored with a warning)
  row 4        : export marker; only columns marked "data" are exported
  row 5+       : rows; the first empty key cell ends the sheet

Enum values for type<ENUM> come from a sheet named "types" (row 2 = enum names,
rows 3+ = members in order) or, if absent, from the distinct cell values in order of appearance.

Usage:
  python -I export_tables.py --xlsx <dir|file> --json-out <dir> --cs-out <dir> [--namespace NS]
  python -I export_tables.py --write-sample <path.xlsx>
"""
import argparse
import json
import keyword
import os
import re
import sys

try:
    import openpyxl
except ImportError:  # pragma: no cover
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    sys.exit(2)

SCALAR_TYPES = {"int", "long", "float", "double", "bool", "string"}
IGNORED_OPTIONS = ("check<", "localtext", "stream", "origin", "null")
CS_RESERVED = {"abstract", "as", "base", "bool", "break", "byte", "case", "catch", "char", "checked", "class", "const",
               "continue", "decimal", "default", "delegate", "do", "double", "else", "enum", "event", "explicit", "extern",
               "false", "finally", "fixed", "float", "for", "foreach", "goto", "if", "implicit", "in", "int", "interface",
               "internal", "is", "lock", "long", "namespace", "new", "null", "object", "operator", "out", "override",
               "params", "private", "protected", "public", "readonly", "ref", "return", "sbyte", "sealed", "short",
               "sizeof", "stackalloc", "static", "string", "struct", "switch", "this", "throw", "true", "try", "typeof",
               "uint", "ulong", "unchecked", "unsafe", "ushort", "using", "virtual", "void", "volatile", "while"}
IDENT = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


class ExportError(Exception):
    pass


class Column:
    def __init__(self, index, name):
        self.index = index
        self.name = name
        self.is_array = name.endswith("[]")
        if self.is_array:
            self.name = name[:-2]
        self.type = None
        self.enum = None
        self.is_key = False
        self.unique = False

    def cs_field(self):
        return "@" + self.name if self.name in CS_RESERVED or keyword.iskeyword(self.name) else self.name

    def cs_type(self):
        base = self.enum if self.enum else self.type
        return base + "[]" if self.is_array else base


def warn(msg):
    print("warning: " + msg, file=sys.stderr)


def parse_options(column, raw, sheet):
    parts = [p.strip() for p in str(raw).split(",") if p.strip()]
    for part in parts:
        low = part.lower()
        if low == "key":
            column.is_key = True
        elif low in SCALAR_TYPES:
            column.type = low
        elif low == "unique":
            column.unique = True
        elif low.startswith("type<") and low.endswith(">"):
            column.enum = part[5:-1].strip()
            column.type = "enum"
        elif low.startswith("typeref<") and low.endswith(">"):
            column.enum = part[8:-1].strip()
            column.type = "enum"
        elif low.startswith(IGNORED_OPTIONS):
            warn(f"{sheet}.{column.name}: option '{part}' is not supported by this exporter; ignored")
        else:
            raise ExportError(f"{sheet}.{column.name}: unknown option '{part}'")
    if column.type is None:
        raise ExportError(f"{sheet}.{column.name}: no type given")
    if column.is_key and column.type not in ("int", "string"):
        raise ExportError(f"{sheet}.{column.name}: key must be int or string")


def convert(column, cell, enums, sheet, row_number):
    def one(text):
        text = "" if text is None else str(text).strip()
        t = column.type
        try:
            if t == "int" or t == "long":
                return int(float(text)) if text else 0
            if t == "float" or t == "double":
                return float(text) if text else 0.0
            if t == "bool":
                return text.lower() in ("1", "true", "yes", "y", "o")
            if t == "string":
                return text
            if t == "enum":
                members = enums.setdefault(column.enum, [])
                if text == "":
                    return 0
                if text not in members:
                    members.append(text)
                return members.index(text)
        except ValueError:
            raise ExportError(f"{sheet}!{column.name} row {row_number}: cannot convert '{text}' to {t}")
        raise ExportError(f"{sheet}.{column.name}: unsupported type {t}")

    if column.is_array:
        text = "" if cell is None else str(cell)
        return [one(p) for p in text.split(",") if p.strip() != ""] if text.strip() else []
    return one(cell)


def read_enum_sheet(workbook, enums):
    if "types" not in workbook.sheetnames:
        return
    ws = workbook["types"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if not rows:
        return
    names = rows[0]
    for col, name in enumerate(names):
        if name is None or str(name).strip() == "":
            continue
        members = []
        for r in rows[1:]:
            v = r[col] if col < len(r) else None
            if v is None or str(v).strip() == "":
                break
            members.append(str(v).strip())
        enums[str(name).strip()] = members


def export_sheet(ws, enums):
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 4 or rows[1][0] is None or str(rows[1][0]).strip().lower() != "#data":
        return None
    sheet = ws.title
    columns = []
    for index in range(1, len(rows[1])):
        name = rows[1][index]
        if name is None or str(name).strip() == "":
            continue
        marker = rows[3][index] if index < len(rows[3]) else None
        if marker is None or str(marker).strip().lower() != "data":
            continue
        column = Column(index, str(name).strip())
        if not IDENT.match(column.name):
            raise ExportError(f"{sheet}: column '{column.name}' is not a valid C# identifier")
        parse_options(column, rows[2][index] if index < len(rows[2]) else "", sheet)
        columns.append(column)

    keys = [c for c in columns if c.is_key]
    if len(keys) != 1:
        raise ExportError(f"{sheet}: exactly one 'key' column is required (found {len(keys)})")
    key = keys[0]

    records = []
    seen_keys = set()
    seen_unique = {c.name: set() for c in columns if c.unique}
    for row_number, row in enumerate(rows[4:], start=5):
        raw_key = row[key.index] if key.index < len(row) else None
        if raw_key is None or str(raw_key).strip() == "":
            break
        record = {}
        for column in columns:
            cell = row[column.index] if column.index < len(row) else None
            record[column.name] = convert(column, cell, enums, sheet, row_number)
        k = record[key.name]
        if k in seen_keys:
            raise ExportError(f"{sheet} row {row_number}: duplicate key '{k}'")
        seen_keys.add(k)
        for name, values in seen_unique.items():
            if record[name] in values:
                raise ExportError(f"{sheet} row {row_number}: duplicate unique value '{record[name]}' in {name}")
            values.add(record[name])
        records.append(record)

    return {"sheet": sheet, "columns": columns, "key": key, "records": records}


def cs_row_class(table, namespace):
    name = table["sheet"]
    key_type = table["key"].type
    lines = [
        "// <auto-generated> tools/tables/export_tables.py; do not edit by hand.",
        "using System;",
        "using BK.Data;",
        "using UnityEngine;",
        "",
        f"namespace {namespace}",
        "{",
        "    [Serializable]",
        f"    public sealed class {name}Row : ITableRow<{key_type}>",
        "    {",
    ]
    for c in table["columns"]:
        lines.append(f"        public {c.cs_type()} {c.cs_field()};")
    lines.append(f"        public {key_type} Id => {table['key'].cs_field()};")
    lines.append("    }")
    lines.append("")
    lines.append(f"    [CreateAssetMenu(menuName = \"BK/Tables/{name}\")]")
    lines.append(f"    public sealed class {name}Table : TableAsset<{key_type}, {name}Row> {{ }}")
    lines.append("}")
    return "\n".join(lines) + "\n"


def cs_enums(enums, namespace):
    lines = ["// <auto-generated> tools/tables/export_tables.py; do not edit by hand.", "", f"namespace {namespace}", "{"]
    for name, members in enums.items():
        lines.append(f"    public enum {name}")
        lines.append("    {")
        for i, m in enumerate(members):
            lines.append(f"        {m} = {i},")
        lines.append("    }")
        lines.append("")
    if lines[-1] == "":
        lines.pop()
    lines.append("}")
    return "\n".join(lines) + "\n"


def export(xlsx_paths, json_out, cs_out, namespace):
    enums = {}
    tables = []
    for path in xlsx_paths:
        workbook = openpyxl.load_workbook(path, data_only=True)
        read_enum_sheet(workbook, enums)
        for ws in workbook.worksheets:
            if ws.title == "types":
                continue
            table = export_sheet(ws, enums)
            if table is not None:
                tables.append(table)
    if not tables:
        raise ExportError("no '#data' sheets found")

    os.makedirs(json_out, exist_ok=True)
    os.makedirs(cs_out, exist_ok=True)
    written = []
    for table in tables:
        json_path = os.path.join(json_out, table["sheet"] + ".json")
        with open(json_path, "w", encoding="utf-8", newline="\n") as f:
            json.dump({"_rows": table["records"]}, f, ensure_ascii=False, indent=2)
            f.write("\n")
        cs_path = os.path.join(cs_out, table["sheet"] + "Table.cs")
        with open(cs_path, "w", encoding="utf-8", newline="\n") as f:
            f.write(cs_row_class(table, namespace))
        written.extend([json_path, cs_path])
    if enums:
        enum_path = os.path.join(cs_out, "TableEnums.cs")
        with open(enum_path, "w", encoding="utf-8", newline="\n") as f:
            f.write(cs_enums(enums, namespace))
        written.append(enum_path)
    return written


def write_sample(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stage"
    ws.append([None])
    ws.append(["#data", "id", "reward", "tags[]", "difficulty", "note"])
    ws.append([None, "key,int", "int", "string", "type<DIFFICULTY>", "string"])
    ws.append([None, "data", "data", "data", "data", "skip"])
    ws.append([None, 1, 100, "a,b", "HARD", "ignored"])
    ws.append([None, 2, 250, "", "EASY", "ignored"])
    types = wb.create_sheet("types")
    types.append([None])
    types.append(["DIFFICULTY"])
    types.append(["EASY"])
    types.append(["NORMAL"])
    types.append(["HARD"])
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--xlsx", help="xlsx file or folder")
    parser.add_argument("--json-out", help="output folder for <Sheet>.json")
    parser.add_argument("--cs-out", help="output folder for <Sheet>Table.cs and TableEnums.cs")
    parser.add_argument("--namespace", default="Project.Tables")
    parser.add_argument("--write-sample", help="write a sample workbook and exit")
    args = parser.parse_args(argv)

    if args.write_sample:
        write_sample(args.write_sample)
        print("sample written: " + args.write_sample)
        return 0
    if not (args.xlsx and args.json_out and args.cs_out):
        parser.error("--xlsx, --json-out and --cs-out are required")

    if os.path.isdir(args.xlsx):
        paths = sorted(os.path.join(args.xlsx, n) for n in os.listdir(args.xlsx)
                       if n.lower().endswith((".xlsx", ".xlsm")) and not n.startswith("~$"))
    else:
        paths = [args.xlsx]
    try:
        for path in export(paths, args.json_out, args.cs_out, args.namespace):
            print("wrote " + path)
    except ExportError as error:
        print("error: " + str(error), file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
